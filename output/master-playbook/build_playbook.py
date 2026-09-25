from pathlib import Path
import json, re, io, base64, struct, uuid, zipfile, html, shutil
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).parent
QA = OUT / 'qa'
QA.mkdir(exist_ok=True)
FONTROOT = ROOT / 'output/website/scent-journey-theme/source-assets/fonts'
FONTS = {
 'Cenzo Flare': FONTROOT/'CenzoFlare-Bold/CenzoFlare-Bold.ttf',
 'Helvetica Now Display': FONTROOT/'HelveticaNowDisplay-Light/HelveticaNowDisplay-Light.ttf',
 'Helvetica Now Text': FONTROOT/'HelveticaNowText-Regular/HelveticaNowText-Regular.ttf',
}
GALLERY = 'output/website/upload-ready/shopify-gallery-order-2026-09-24/'
THEME = 'output/website/scent-journey-theme/'
REF = 'handoff/assets/brand-references/'
PAGEKIT = 'output/facebook-page-kit/priority-six-2026-09-25-v01/'

def asset(slug, prefix):
    return str(next((ROOT/GALLERY/slug).glob(prefix+'*.png')).relative_to(ROOT)).replace('\\','/')

pages=[]
def page(title, section, source=''):
    d={'title':title,'section':section,'source':source,'blocks':[]}; pages.append(d); return d['blocks']
def p(b, text, style='Normal'): b.append({'kind':'p','text':text,'style':style})
def h(b,text): p(b,text,'Heading 2')
def bullets(b, items):
    for x in items: p(b,x,'List Bullet')
def table(b, headers, rows, widths=None): b.append({'kind':'table','headers':headers,'rows':rows,'widths':widths})
def pic(b,path,caption,width=6.55): b.append({'kind':'image','path':path,'caption':caption,'width':width})
def rowpics(b,paths,captions,width=3.17): b.append({'kind':'images','paths':paths,'captions':captions,'width':width})

b=page('Dear Body Master Playbook','BRAND FOUNDATION TO EVERYDAY EXECUTION','Version 1.0 | 25 September 2026')
p(b,'Brand story  Visual identity  Products  Content  Website  Operations  AI handoff','Subtitle')
pic(b,THEME+'theme/assets/sj-logo-official-dark.png','Official Dear Body identity',3.8)
rowpics(b,[asset('mojito-metallique','07-'),asset('charme-envoutant','07-')],['Warm color, tactile objects and an unmistakable product.','A distinct scent world, expressed through hands and detail.'],3.17)
p(b,'A shared guide for everyone connected to DearBody Philippines: the owner, collaborators, designers, photographers, copywriters, developers, customer care teams and AI platforms.')
p(b,'The brand system brings London-formulated fragrance culture into Filipino everyday life through warmth, color, personal style and approachable product storytelling. This playbook connects that idea to the exact decisions and assets used to build it.')

b=page('How to use this playbook','START HERE','S01–S04 and S11')
p(b,'Use this document to understand the brand, brief work, create consistent assets and check the result. Read the quick rules first, then the section relevant to your task. The final chapters contain source locations and a reusable AI handoff.')
h(b,'The contents')
b.append({'kind':'toc'})
p(b,'Decision language: “Current rule” records an owner direction or established production standard. “Recorded implementation” describes the project at a stated date. “Recommended workflow” is practical guidance for future work. “Historical” material explains earlier development and does not override current rules.','Caption')

b=page('The brand in one page','01  BRAND FOUNDATIONS','S01 Brand guide | S02 Owner rules | S03 Brand tokens')
p(b,'DearBody Philippines is premium but approachable: London-formulated fragrance culture translated into warm, vibrant, playful Filipino self-expression. The central idea is A Scent Journey.')
h(b,'The rules everyone must know')
bullets(b,[
 'Use the real Dear Body identity. Preserve official logo artwork and exact product packaging, names, label colors, cap construction and scale.',
 'Lead with deep burgundy, rich red, burnt orange, golden tan and soft cream. Scent colors add individuality; warm hues anchor the brand.',
 'Use actual Cenzo Flare Bold for display copy and Helvetica Now for support. The website uses Helvetica Now Display Light; earlier Text Regular assets remain archival.',
 'Current Facebook scope is six ads, one per priority scent, plus one profile image and one cover. A fourteen-scent archive does not expand this campaign.',
 'Current campaign lifestyle and scripts are face-free: Filipino hands, arms and torso crops only. Exclude partial faces, reflections, background people, screens and printed portraits.',
 'Women’s white or cream labels: Mojito Metallique, Amber Oud Silk, Mistened Narcissus. Men’s black labels: Charme Envoûtant, Oud Mirage, Rtulle & Satin. Match all visible talent to the line.',
 'Offer order is FREE SHIPPING ON 2+ ITEMS, then CASH ON DELIVERY, then SHOP NOW. Keep the condition next to the shipping message.',
 'Use approved scent descriptions. Do not invent notes, performance, origin claims, ingredients, discounts, popularity or testimonials.',
 'Shopify is the master sales channel in the current project. Preserve real product data and the recorded EasySell order flow.',
 'New owner decisions override older instructions. Existing product-gallery artwork stays unchanged unless an edit is requested.'
])

b=page('How the brand system was built','01  BRAND FOUNDATIONS','S01–S05 | S08–S12 | S15 Historical research')
p(b,'The project documents the development of the DearBody Philippines brand experience around existing DearBody products, packaging and official marks. It records local strategy, creative direction, production systems and commerce. This history should not be retold as the invention of the global fragrance house or its formulas.')
table(b,['Stage','What was established','What it enabled'],[
 ['Early strategy archive','An accessible premium fragrance proposition; personality, bold visuals and online discovery.','A practical reason for buyers to notice and identify with the brand.'],
 ['22 September 2026','The 12-page Canva playbook was translated into BRAND_GUIDE.md, brand tokens and visual rules.','A shared palette, typography system, mission and tone.'],
 ['Product evidence and galleries','Owner photographs, verified package facts, exact scent descriptions and a twelve-frame image architecture.','Repeatable product storytelling with packaging fidelity.'],
 ['24 September 2026','Fourteen gallery sets were reviewed; six selected scents were integrated into the website.','A broader asset library with a focused active storefront.'],
 ['24 September 2026','The Scent Journey Shopify theme, display-font update, pricing corrections and same-page COD ordering were recorded.','A connected visual and purchasing experience.'],
 ['25 September 2026','Owner references sharpened the social direction toward energetic, warm, tactile, face-free campaign imagery.','A clear creative standard for the next six-scent Facebook campaign.']
 ],[1.25,2.95,2.35])
h(b,'The idea that stayed consistent')
p(b,'Fragrance belongs in ordinary life and gives people another way to express themselves. The visual execution became more specific over time: richer color, stronger product identity, real texture, varied activities and clear buying information.')

b=page('Positioning and audience','01  BRAND FOUNDATIONS','S01 Brand guide | S15 Historical strategy')
h(b,'Positioning')
p(b,'London-formulated fragrance culture curated for the Filipino lifestyle: premium, warm, vibrant, playful and approachable. Use “London-formulated” as the established brand wording. Do not silently upgrade it to “made in London,” “British-made,” or a manufacturing claim.')
h(b,'Mission and vision')
p(b,'Mission: To bring carefully curated fragrances from London to the Philippines, creating scents that complement Filipino tastes, lifestyle and the tropical environment.')
p(b,'Vision: To bring London’s fragrance culture into every Filipino’s everyday life.')
h(b,'What the audience is looking for')
p(b,'The early questionnaire describes online shoppers aged 22–35, including students and working individuals, who care about personal image and distinctive scents. Treat this as a planning persona, not a restriction on who can enjoy the products or a verified description of every customer.')
p(b,'The practical need is affordable, effortless everyday confidence. The emotional need is a scent that feels personal. Help shoppers recognize a mood, understand the product and choose without needing specialist fragrance vocabulary.')
h(b,'How to express the position')
bullets(b,['Show appealing packaging and believable personal objects. Let the quality appear in the photograph and the clarity of the experience.','Use friendly confidence and short sensory language. Make the product feel attainable in a real Filipino routine.','Build differentiation through consistent design and distinct scent worlds. Avoid unsupported claims that competitors lack branding or that Dear Body is the only brand in a category.'])

b=page('Brand voice and story','01  BRAND FOUNDATIONS','S01 | S06 Approved copy | S09 Theme story section')
h(b,'Core brand story')
p(b,'DearBody Philippines brings London’s fragrance culture into Filipino everyday life through London-formulated scents thoughtfully curated for local tastes, lifestyles, and tropical conditions.')
p(b,'We believe fragrance belongs in the everyday: the morning you make your own, the plans you say yes to, the moments that become memories. International fragrance character meets Filipino personality — and a little more room for self-expression.')
p(b,'The two paragraphs above are the recorded website story, suitable as the base for introductions. Keep a personal founder biography separate until the owner supplies the people, dates and experiences behind it.','Caption')
table(b,['Quality','How the copy behaves'],[
 ['Warm','Talk to one person in familiar words. Invite exploration without pressure.'],['Confident','Use a direct headline and a clear next action. Avoid exaggerated superiority.'],['Playful','Let the setting, gesture or phrase carry a little energy. Avoid forced slang.'],['Sensory','Use the approved description. Do not turn evocative wording into a laboratory claim.'],['Approachable','Use English for clarity; natural Filipino or Taglish can support social scripts. Keep offer conditions unambiguous.']
 ],[1.35,5.2])
h(b,'Reusable language')
p(b,'A SCENT JOURNEY is the brand tagline. MAKE IT PERSONAL. is the recorded homepage headline. “Six fragrances. Plenty of personality. Find the one that feels like you.” is a concise storefront introduction. These serve different roles; keep the tagline consistent.')
p(b,'Example caption structure: offer and condition → exact scent name → approved description → Shop Now. Personality copy is an editorial invitation, not a promise that a fragrance changes someone’s character.')

b=page('The color system','02  VISUAL IDENTITY','S01 Brand guide | S03 Brand tokens')
p(b,'The warm palette is the consistent brand signal across products, publication layouts, social content and the website. Use contrast to keep glass, labels and typography clear.')
b.append({'kind':'palette'})
table(b,['Color','HEX','RGB','Role'],[
 ['Deep burgundy','#5C0006','92  0  6','Primary anchor; premium dark fields'],['Rich red','#9A1106','154  17  6','Buttons, emphasis, active states'],['Burnt orange','#D46601','212  102  1','Energy and campaign accents'],['Golden tan','#E9A250','233  162  80','Warm fields and secondary color'],['Soft cream','#F4E3CB','244  227  203','Light fields and contrast']
 ],[1.5,1.05,1.25,2.75])
h(b,'Use color with a purpose')
p(b,'Start with one brand anchor, one real scent accent and one neutral. Keep the overall composition warm, even when Oud Mirage uses vivid blue or Rtulle & Satin uses aqua. Saturated accents in the owner references are compatible with the current playful direction; they do not change the official five-color palette.')
p(b,'For small copy on cream, use dark burgundy or near-black. On dark fields, use cream or a sufficiently light tone. Check the final export at mobile size. Do not rely on pale orange, yellow or lavender text over cream.')
p(b,'The HEX values are digital specifications. Obtain printer-specific proofs before treating any CMYK conversion as an approved print match. Scent accent colors should be sampled from the approved source product, not invented as fixed HEX values.')

b=page('Typography and logo rules','02  VISUAL IDENTITY','S01 | S03 | S10 Supplied font notes | Official logo assets')
p(b,'Cenzo Flare Bold','Specimen Display')
p(b,'A SCENT JOURNEY','Specimen Display')
p(b,'Helvetica Now Display Light','Specimen Body')
p(b,'Warm, playful fragrance for everyday self-expression.','Specimen Body')
table(b,['Application','Production face','Use'],[
 ['Display','Cenzo Flare Bold','Headlines, campaign titles and major statements'],['Live website support','Helvetica Now Display Light 300','Body copy, navigation, UI and captions'],['Earlier gallery typesetting','Helvetica Now Text Regular 400','Preserved source for archived exports; do not relabel old raster artwork'],['Editable playbook body','Helvetica Now Text Regular','Long-form reading; display examples retain the actual brand faces']
 ],[1.65,2.2,2.7])
rowpics(b,[THEME+'theme/assets/sj-logo-official-dark.png',PAGEKIT+'final/dear-body-facebook-profile-1080.png'],['Official primary wordmark. Preserve the artwork and proportions.','Existing PH emblem profile export. Review its legibility in a small circle.'],2.2)
p(b,'Use the wordmark for core brand placements and the circular monogram for small placements. Protect contrast and leave generous breathing room. Do not redraw, stretch, recolor arbitrarily or add a PH suffix to the official global mark. Use an approved PH export when needed.')
p(b,'The supplied fonts are embedded in this Word file. Keep the original licensed files with the production handoff. Technical embedding permission is not a transfer of font-license ownership. Final advertising requires actual brand fonts; a generic fallback must be flagged for approval.','Caption')

b=page('The owner visual references','02  VISUAL IDENTITY','S02 Owner direction dated 25 September 2026 | S04 Reference images')
pic(b,REF+'2026-09-25-publication-materials-owner-reference.png','Publication reference: tight editorial crops, saturated objects, tactile clothing and product-led personal style.',6.55)
pic(b,REF+'2026-09-25-social-media-look-owner-reference.png','Social reference: warm hues dominate a varied, energetic feed.',6.55)
p(b,'Read these as visual-language references. They include historical faces and third-party inspiration; the current Facebook campaign must use face-free crops. Recreate the energy and material quality, not the exact composition, people or third-party assets.')

b=page('Photography and packaging fidelity','02  VISUAL IDENTITY','S02 | S05 Product image strategy | S07 Product source facts')
rowpics(b,[asset('mojito-metallique','07-'),asset('amber-oud-silk','07-')],['Mojito Metallique: warm daylight, chrome, yellow and cobalt.','Amber Oud Silk: realistic bottle scale, canvas and an active packing gesture.'],3.17)
h(b,'What makes a Dear Body image')
p(b,'Build a believable moment around the product: a hand lifting it, a bag being packed, a cap being removed, an object being set down. Use warm daylight or warm editorial light, natural skin texture and physically credible shadows. Give each scent a distinct action and prop world.')
h(b,'What must remain exact')
bullets(b,['Use the owner’s product photography as the packaging authority. Preserve the cylindrical glass, heavy base, real liquid color, logo, label, cap and proportions.','The glossy black cap and the separate silver atomizer collar are different components. Do not add a silver underside to the detached cap or invent extra canister pieces.','Check the label at full size and at mobile size. Prefer the real product cutout when label accuracy cannot be maintained through generation.','Generate or photograph the scene first; add all new headlines and offer copy with actual fonts afterward.'])

SCENTS=[
 ('mojito-metallique','Mojito Metallique','Women','Soft, creamy sweetness with a light fruity glow.','Yellow canister; golden-yellow liquid; white label.','P11138','5056795407345','For someone friendly and easygoing, who puts people at ease and brings gentle humor to everyday life.','FRIENDLY · EASYGOING · LIGHTHEARTED','Dance rehearsal and music','Yellow headphones, cobalt bag, coral towel and warm studio daylight. Keep the camera below the chin and away from mirrors.','Creamy and fruity wording does not authorize mint, lime, cocktails or a fresh-citrus note pyramid.'),
 ('amber-oud-silk','Amber Oud Silk','Women','Dark spice and smoke with a warm, addictive edge.','Blush-peach canister and liquid; white label.','P11038','5056795407338','For someone bold and self-assured, who brings warmth to their relationships and enjoys expressing their individuality.','BOLD · WARM · SELF-ASSURED','Apartment entry and getting ready','Coral cotton, navy canvas, keys and a wooden console. Use a face-free hand lifting or packing the bottle.','The current gallery story is the realistic apartment routine. The former karaoke direction is superseded; smoke is descriptive copy, not a required effect.'),
 ('mistened-narcissus','Mistened Narcissus','Women','Juicy berries softened by a smooth, sweet finish.','Pale lavender canister; clear liquid; white label.','P11438','5056795407376','For someone cheerful and playful, who shows their affection openly and finds joy in little moments.','CHEERFUL · PLAYFUL · OPENHEARTED','Roller-rink preparation','Lavender skates, coral details, chrome and warm rink light. Film seated hands or a bag still life before skating.','Preserve the exact spelling Mistened Narcissus. Do not handle a glass perfume bottle while actively skating.'),
 ('charme-envoutant','Charme Envoûtant','Men','Rich, spiced sweetness that feels deep and indulgent.','Orange canister; amber-orange liquid; black label.','P10738','5056795407307','For someone expressive and passionate, who follows their enthusiasm and enjoys making bold personal choices.','EXPRESSIVE · PASSIONATE · CONFIDENT','Parked convertible and personal style','Orange and red styling, chrome details and a male hand at a parked car. No driving action, face or face in a mirror.','Retain the circumflex in Envoûtant. Clothing and posture should feel expressive and approachable; avoid reusing another scent’s sports story.'),
 ('oud-mirage','Oud Mirage','Men','Dark rose wrapped in smoky, woody depth.','Vivid blue canister and transparent blue liquid; black label.','P10238','5056795407253','For someone independent and reflective, who trusts their own taste and feels comfortable taking a different path.','INDEPENDENT · REFLECTIVE · SELF-ASSURED','Weekend park and bicycle bag','Blue against coral, sunlit grass and a practical bag interaction. Male hands only; avoid background people and reflected faces.','The current lifestyle is relaxed weekend leisure. The working-professional or architecture-studio story is superseded.'),
 ('rtulle-and-satin','Rtulle & Satin','Men','Airy sweetness with a warm, glowing trail.','Turquoise or aqua canister and liquid; black label.','P10038','5056795407239','For someone gentle and imaginative, who sees beauty in small things and shares their warmth with a lighthearted spirit.','GENTLE · IMAGINATIVE · LIGHTHEARTED','Pottery-painting café','Aqua and coral, ceramic forms, a tote and male hands. Keep perfume distinct from cups and art materials.','Use Rtulle & Satin in display copy. The source folder is rtulle-and-satin; the current Shopify handle is rtulle-satin. Do not revive the tailoring or sewing concept from the word Satin.')
]
for slug,name,line,desc,appearance,code,barcode,person,traits,world,scene,caution in SCENTS:
    b=page(name,'03  THE SIX PRIORITY SCENTS','S06 Exact copy ledger | S07 Product facts | S08 Personality review')
    rowpics(b,[asset(slug,'01-'),asset(slug,'07-')],['Product identity and packaging reference.','Existing face-free lifestyle example.'],2.65)
    p(b,desc,'Lead')
    p(b,f'{line}’s line  |  50 ml / 1.69 fl. oz.  |  {code}  |  Barcode {barcode}','Caption')
    p(b,appearance)
    h(b,'Personality interpretation')
    p(b,person)
    p(b,traits,'Caption')
    h(b,world)
    p(b,scene)
    p(b,caution)

b=page('The wider product library','03  THE SIX PRIORITY SCENTS','S06 Scent description ledger | S08 Gallery review')
p(b,'The project contains fourteen completed gallery sets. Six are the priority campaign and storefront range documented here. The remaining eight belong to the wider asset library; their inclusion below is a reference record, not authorization to add products to the campaign or claim current availability.')
ledger=(ROOT/'output/product-listing/SCENT_DESCRIPTION_SOURCE_LEDGER.md').read_text(encoding='utf-8')
others=[]
priority_names=[s[1] for s in SCENTS]
for line in ledger.splitlines():
    if line.startswith('| ') and 'Complete:' in line:
        cells=[x.strip() for x in line.strip('|').split('|')]
        if cells[0] not in priority_names: others.append([cells[0],cells[1].strip('“”')])
table(b,['Additional scent','Approved description in the archive'],others,[1.6,4.95])
h(b,'Expansion rule')
p(b,'Before activating another scent, confirm the requested scope, package facts, label category, approved copy, stock, price, product handle and selected gallery. Do not reuse six-scent pricing assumptions for the wider library. The fourteen-scent upload package is prepared material; it is not evidence that all fourteen are published.')

b=page('Product gallery architecture','04  PRODUCT CONTENT','S05 Production strategy | S08 Current upload order and import QA')
p(b,'Each scent has twelve distinct frames. Keep the production frame identity separate from the Shopify display position. The later upload order puts scent choice information immediately after the main image.')
table(b,['Shopify position','Frame purpose','Production frame'],[
 ['01','Main image with no added offer copy','01'],['02','Exact scent description','11'],['03','Who it fits and three personality traits','12'],['04','Product profile and verified facts','08'],['05','Ingredients and care','09'],['06','Lifestyle hero','05'],['07','Product in hand lifestyle','06'],['08','Editorial lifestyle','07'],['09','Spray in motion','02'],['10','Packaging reveal','03'],['11','Magnetic cap detail','04'],['12','Closing hero','10']
 ],[1.25,3.8,1.5])
p(b,'The recorded website import order is 01, 11, 12, 08, 09, 05, 06, 07, 02, 03, 04, 10. The upload-ready folders are already numbered in display order. Follow their manifest; do not upload them in a second re-sorted production order.')
h(b,'Technical handoff')
p(b,'Approved source galleries are 2048 × 2048 PNG files with a square master and sRGB workflow. The website import keeps unchanged source PNGs plus 1200 × 1200 JPEG derivatives and 160 × 160 thumbnails. Retain the source and checksum record. Inspect readability on the real mobile gallery and use descriptive alt text.')

b=page('Gallery examples and copy layers','04  PRODUCT CONTENT','S05 | S06 | S08')
rowpics(b,[asset('mojito-metallique','02-'),asset('mojito-metallique','04-')],['Scent description: the approved sentence is the source of truth.','Product profile: exact package facts make the product easier to inspect.'],3.17)
rowpics(b,[asset('charme-envoutant','11-'),asset('oud-mirage','10-')],['Construction detail: preserve the real cap and atomizer relationship.','Packaging reveal: preserve the actual canister components.'],3.17)
p(b,'These are embedded examples from the existing gallery library. New campaign artwork follows the newer face-free rule. Keep photographic masters separate from text overlays so copy can be corrected without regenerating the product.','Caption')

b=page('Copy accuracy and claims','04  PRODUCT CONTENT','S05 | S06 | S08 | S15 Historical research')
h(b,'Three different types of product language')
table(b,['Type','Authority','How to use it'],[
 ['Physical fact','Owner photographs and verified package transcription','Volume, code, barcode, ingredients, construction and safety wording must be exact.'],['Scent description','Owner-approved copy ledger','Use the exact sentence for description modules. Do not reverse the Mojito and Amber descriptions.'],['Personality interpretation','24 September personality review','Use the “For someone…” sentence and three traits as editorial wearer language. It is not a product-performance claim.']
 ],[1.35,2.0,3.2])
h(b,'Copy that requires separate evidence')
p(b,'Longevity in hours, projection, oil percentage, hypoallergenic claims, “safe for sensitive skin,” ingredient benefits, manufacturing country, certifications, universal suitability and guaranteed compliments require product-specific approval and evidence. The name Oud, Mojito or Narcissus is not a license to invent a note pyramid.')
h(b,'Ingredients are not fragrance notes')
p(b,'Transcribe each product’s ingredient panel separately. Do not copy one scent’s allergen list to another, infer a top-heart-base structure, or turn an ingredient into a benefit claim. Keep the approved full ingredient panel and safety text with the product brief.')
h(b,'Historical material that must not become live copy')
p(b,'The early strategy PDF includes suggested longevity language, vouchers, bundle prices, competitor comparisons and estimated market data. Those suggestions do not establish approved product claims or current offers. The later package-fidelity and owner-copy standards control production.')
h(b,'Useful caption example')
p(b,'Free shipping on 2+ items + cash on delivery. Meet Mojito Metallique: soft, creamy sweetness with a light fruity glow. Choose your scents and tap Shop Now. #DearBodyPH #AScentJourney #MojitoMetallique')

b=page('The current Facebook campaign','05  SOCIAL AND VIDEO','S02 Owner rules | S04 Visual references | S12 Existing social assets')
p(b,'Current scope: six ads total, one for each priority scent, plus a profile image and a cover. The ad system needs a shared brand identity and a distinct scene for every fragrance.')
h(b,'Offer hierarchy')
p(b,'FREE SHIPPING ON 2+ ITEMS','Specimen Display')
p(b,'CASH ON DELIVERY','Specimen Body')
p(b,'SHOP NOW','Specimen Display')
p(b,'Make the two-item condition part of the first message. Support it with COD, then the exact scent name and a clear action. Keep the product label uncovered. Do not add a sale deadline, voucher, extra discount or geographic promise without confirming the offer.')
h(b,'Creative construction')
bullets(b,['Use warm, saturated color, tactile photography, tight crops and energetic personal objects. Avoid repetitive plain banner panels and generic softly lit smiling portraits.','Show Filipino hands, arms and torso only. Review mirrors, chrome, glass, shadows, screens, prints and background figures for faces or face fragments.','Use female talent for the three cream-label scents and male talent for the three black-label scents. The category rule applies to every visible body or hand.','Add live type or deterministic typography after the photographic base. Use actual Cenzo Flare Bold and Helvetica Now.'])
h(b,'Existing output sizes')
p(b,'The earlier ad exports are 1080 × 1350 at 4:5. The existing profile export is 1080 × 1080, and the existing cover is 1640 × 720. These are project output sizes; verify the intended platform placement and crops before final publication.')
p(b,'Status: the recorded six-ad exports and existing cover show faces and need revision before reuse under the 25 September direction. Their file dates alone do not make them compliant. The existing profile is a brand-identity reference. This playbook does not publish or launch ads.','Caption')

b=page('Six distinct content concepts','05  SOCIAL AND VIDEO','S02 current rule | S07 scent worlds | Adapted production concepts')
p(b,'Recommended treatments below adapt the established scent worlds to the current face-free direction. The shared offer stays constant; the gestures, objects, wardrobe and environment create variation.')
table(b,['Scent','Face-free action and visual hook','What to avoid'],[
 ['Mojito Metallique','A woman’s hand lifts the bottle from beside yellow headphones and a cobalt dance bag.','Mirrors, a visible chin, mint or lime props inferred from the name.'],
 ['Amber Oud Silk','A woman packs the bottle into a navy canvas bag on a wooden entry console.','Karaoke revival, oversized bottle, duplicate keys or bags.'],
 ['Mistened Narcissus','A seated woman sets the bottle beside lavender skates, then reaches for her bag.','Glass handling during skating; rink crowds and reflected faces.'],
 ['Charme Envoûtant','A male hand and orange sleeve lift the bottle beside a parked convertible’s chrome detail.','Driving, face reflections in car mirrors, duplicated sunglasses.'],
 ['Oud Mirage','A male hand brings the blue bottle out of a bicycle bag at a sunlit park bench.','Architecture-office story, background walkers or cold brooding light.'],
 ['Rtulle & Satin','A man’s cropped hands arrange the aqua bottle beside a tote and painted ceramics.','Tailoring literalization; perfume placed where it resembles a drink or paint.']
 ],[1.4,3.15,2.0])
h(b,'A varied feed')
p(b,'Rotate clear product identity, scent-description copy, face-free personal moments, packaging detail, a shipping explanation and an invitation to explore the Scent Finder. Keep recurring motifs recognizable without repeating one pose or composition across six scents.')
p(b,'Recommended test: change one opening action or crop while keeping the offer, scent, audience and destination consistent. Compare attention and completed orders separately. A high view count does not by itself prove purchasing performance.')

b=page('Short video scripts for women','05  SOCIAL AND VIDEO','Recommended scripts based on S02 and S06')
p(b,'Use a 15-second vertical edit. Product identity appears immediately. All shots are cropped below the chin, with no faces anywhere. Use licensed music, real cloth or object sounds, and readable captions.')
for slug,name,line,desc,appearance,code,barcode,person,traits,world,scene,caution in SCENTS[:3]:
    h(b,name)
    action={'mojito-metallique':'Hand lifts the bottle beside headphones; crop at the torso.','amber-oud-silk':'Hand lifts the bottle from the console; navy bag stays in frame.','mistened-narcissus':'Bottle rests beside skates; seated hand turns its label to camera.'}[slug]
    table(b,['Time','Picture','Spoken or screen message'],[
      ['0–4 sec',action,'“Free shipping kapag two items or more!”\nFREE SHIPPING ON 2+ ITEMS'],
      ['4–6 sec','Hold the product steady; make one small prop interaction.','“May cash on delivery din.”\nCASH ON DELIVERY'],
      ['6–11 sec','Readable label close-up and the exact scent name.',desc],
      ['11–15 sec','Clean product hold or still life with an uncluttered end card.','“Choose your scents and tap Shop Now.”\nSHOP NOW']
    ],[0.8,2.75,3.0])
p(b,'Keep the description as a readable on-screen line when the voiceover would be rushed. Do not squeeze all four messages into one crowded opening frame.','Caption')

b=page('Short video scripts for men','05  SOCIAL AND VIDEO','Recommended scripts based on S02 and S06')
p(b,'Use male hands and torso crops for these black-label scents. Keep skin, grip and scale natural. Record the label close-up with a steady camera and an unobstructed front panel.')
for slug,name,line,desc,appearance,code,barcode,person,traits,world,scene,caution in SCENTS[3:]:
    h(b,name)
    action={'charme-envoutant':'Male hand lifts the bottle beside a parked car’s chrome; exclude mirrors.','oud-mirage':'Male hand opens a bicycle bag at a bench and reveals the blue bottle.','rtulle-and-satin':'Male hand sets the aqua bottle beside ceramics and a canvas tote.'}[slug]
    table(b,['Time','Picture','Spoken or screen message'],[
      ['0–4 sec',action,'“Free shipping kapag two items or more!”\nFREE SHIPPING ON 2+ ITEMS'],
      ['4–6 sec','Brief object detail with the product still visible.','“May cash on delivery din.”\nCASH ON DELIVERY'],
      ['6–11 sec','Label close-up; exact scent name in live type.',desc],
      ['11–15 sec','Static product end card; clear space around the label.','“Choose your scents and tap Shop Now.”\nSHOP NOW']
    ],[0.8,2.75,3.0])
p(b,'For a 30-second version, extend the action and product-detail beats. Do not fill the added time with unsupported performance claims, extra scenes or a visible portrait.','Caption')

b=page('The Shopify website system','06  WEBSITE AND COMMERCE','S09 Current theme START-HERE and source')
p(b,'The current project website is the Scent Journey theme for dearbody.ph. It is a native Shopify Online Store 2.0 implementation using Liquid, JSON templates, sections, snippets, CSS and vanilla JavaScript. Earlier theme folders are retained design history.')
table(b,['Layer','Purpose'],[
 ['Brand and navigation','Official identity, search and optional Light or Dark switch.'],['Homepage','A product-free MAKE IT PERSONAL. hero; three feature photos; Women and Men campaign gateways.'],['Collection pages','Browse the women’s or men’s scents, select items and set order quantities.'],['Product pages','Twelve-frame gallery, approved story, native product data and order controls.'],['Scent Finder','Five-question discovery flow returning one scent to explore.'],['Customer information','Our Story, FAQ, Shipping and Contact; native policy content where available.'],['Ordering','Same-page EasySell COD flow with genuine saved-order confirmation.']
 ],[1.45,5.1])
h(b,'Two art-directed modes')
p(b,'Light and Dark share the catalog and commerce behavior. The dark version is a considered visual direction, not a simple color inversion. The sun or moon switch may be enabled in theme settings and retains the choice while browsing in the same tab.')
h(b,'What should stay native')
p(b,'Titles, variants, IDs, stock, selling prices and availability come from Shopify. A theme ZIP cannot create inventory, shipping rules, external app settings or legal policies. The local preview is a review environment and cannot place orders.')

b=page('Website visual examples','06  WEBSITE AND COMMERCE','S09 Archived local QA screenshots and theme assets')
pic(b,THEME+'qa/screenshots/home-hero/home-light-1440-top.png','Recorded Light website layout from the home-hero revision. It illustrates hierarchy and navigation; it is an earlier visual snapshot, not a fresh capture of the live store.',6.55)
rowpics(b,[THEME+'qa/screenshots/gallery-controls/behavior-1440-light.png',THEME+'qa/screenshots/gallery-controls/layout-1440-dark.png'],['Product gallery layout in Light.','Product gallery layout in Dark.'],2.75)
p(b,'Screenshots preserve their tested revision. Their visible faces and any older controls or copy are historical evidence; use the written current rules for new social work and the current theme source for development. Existing website and gallery images are not automatically replaced by the new Facebook direction.','Caption')

b=page('Website pages and interaction rules','06  WEBSITE AND COMMERCE','S09 START-HERE | Source templates | Local QA records')
h(b,'Homepage')
p(b,'The current documented homepage has no individual-product shopping grid. The opening hero stays product-free and its “Explore the collections” action scrolls to the Women and Men gateways. Three photographic features explain the magnetic cap, collectible bottle display and free shipping on two or more items. Those feature sections have no buttons.')
h(b,'Collections and product pages')
p(b,'Women and Men show the actual product selection. On mobile, keep every collection card before the order form. On product pages, maintain a usable gallery, readable product copy and selected quantities. Do not stack twelve full-size images into a long mobile page. Preserve zoom, keyboard controls, reduced-motion behavior and the recorded gallery order.')
h(b,'Banners and typography')
p(b,'The recorded unified banner height is clamp(540px, 43vw, 650px) on desktop and 600px on mobile. Hero text remains aligned to shared content gutters, even when photography spans the viewport. Mobile copy overlays the photographs. Category gateways retain their portrait composition. Test each edit across narrow and wide screens.')
h(b,'Native resource mapping')
table(b,['Page handle','Template'],[['our-story','page.our-story'],['scent-finder','page.scent-finder'],['faq','page.faq'],['contact','page.contact'],['shipping','page.shipping'],['privacy and terms','Native policy content or corresponding page template when populated']],[2.0,4.55])
p(b,'The recorded website removed Returns links, the returns-related FAQ content and the Returns template. Route actual return concerns to the owner’s current approved policy; do not invent a policy.')

b=page('Scent Finder and search content','06  WEBSITE AND COMMERCE','S09 START-HERE | sj-finder-profile.liquid | MASTER-BRIEF source copy')
h(b,'What the Finder does')
p(b,'Five questions cover collection, mood, scent style, occasion and scent feel. Collection filters the six candidates first. The documented scoring then gives scent style 5 points, mood 3, occasion 1 and scent feel 1. Show the highest-scoring scent only; ties retain catalog order.')
p(b,'If every preference is skipped, show one clearly labeled suggestion to explore rather than claiming a personalized match. Sold-out products keep their true availability unless “Available products only” is enabled. An empty eligible collection should produce an honest empty state.')
h(b,'What the Finder does not establish')
p(b,'Mood, occasion and intensity fields are editorial discovery suggestions. “Rich and deep” or “soft and airy” describes character, not measured longevity or projection. The Finder’s shortened descriptions are implementation copy; use the exact owner sentence in the formal scent-description module.')
h(b,'Editable product fields')
table(b,['Namespace custom','Purpose'],[['scent_finder_enabled','Explicit false excludes a product.'],['mood and scent_character','Override editorial matching traits.'],['occasion and intensity','Override discovery suggestions, without implying tested performance.'],['hero_statement and short_description','Approved introductions for cards or product pages.'],['scent_description and personality','Approved long copy and wearer interpretation.']],[2.5,4.05])
h(b,'Search and SEO')
p(b,'Use a clear title containing DearBody, the exact scent name and an approved character description. Example: “DearBody Mojito Metallique Perfume | Soft, Creamy & Fruity Fragrance.” Keep product metadata consistent with the actual page, add descriptive image alt text and avoid unsupported “long-lasting” keywords.')

b=page('Prices offers and ordering','06  WEBSITE AND COMMERCE','S09 START-HERE | S11 COD update | Recorded 24 September 2026')
p(b,'The following is the recorded commercial configuration, not a fresh stock or pricing check. Before publishing a new offer, verify native Shopify values and the applicable shipping and app settings.')
table(b,['Recorded item','Value or rule'],[['Six active selling prices','₱799 each'],['Owner-approved original reference','₱1,099 for the six eligible active scents'],['Displayed savings','₱300 or 27.3% OFF; (1099 − 799) ÷ 1099 × 100'],['One item in the Philippines','₱80 shipping; recorded total ₱879 at ₱799'],['Two or more items','Free shipping, including two bottles of the same scent'],['Two bottles at recorded prices','₱1,598 with free shipping'],['Three bottles at recorded prices','₱2,397 with free shipping']],[2.3,4.25])
h(b,'The customer journey')
p(b,'Choose Add to order → select scents and quantities → choose Check out → enter delivery details once in the EasySell popup → submit → receive genuine order confirmation after the order is saved. An opened popup is not a placed order.')
p(b,'Keep exact native variant IDs and quantities, merge duplicate variants correctly and validate the prepared cart before opening the app. The product page supports its main scent and an optional second scent. Women’s and men’s selections can share an order.')
h(b,'Pricing details for maintainers')
p(b,'The record states that only Mojito had the native ₱1,099 compare-at value; the other five used the owner-approved theme reference. Other currencies, mixed prices and non-priority products retain native compare-at behavior. Do not apply another discount at checkout or revive the earlier 20% or ₱998.75 reference.')

b=page('Delivery customer care and store operations','06  WEBSITE AND COMMERCE','S09 Shipping content | S11 COD and app evidence')
table(b,['Destination','Recorded delivery estimate'],[['Luzon','2–3 days'],['Visayas','3–5 days'],['Mindanao','5–10 days']],[2.1,4.45])
p(b,'Present these as the owner’s recorded estimates, not guaranteed arrival dates. Confirm the current courier and dispatch conditions when answering an individual order inquiry.')
h(b,'Approved factual reply patterns')
bullets(b,['Shipping: “The recorded shipping fee is ₱80 for one item, and orders with two or more items receive free shipping in the Philippines.” Recheck the current offer before using this as a live reply.','Choosing a scent: quote the appropriate approved sentence, offer the Scent Finder and explain differences without promising the customer will love a scent.','Order status: check the actual order record. Ask for the minimum order-identifying information through the approved support channel; do not paste customer details into a public AI prompt.','Performance questions: share approved product information. Do not invent a wear time, oil percentage or skin-safety guarantee.'])
h(b,'EasySell operating record')
p(b,'The 24 September release records the Free plan with 60 orders per month and usage observed at 1 of 60 after testing. Paid SMS and WhatsApp services were deactivated. These are dated settings, not a current usage reading. No paid upgrade or messaging purchase is implied by this playbook.')
p(b,'The genuine same-page COD test was completed, then canceled, voided and archived with three bottles restocked. The subsequent branded confirmation appearance was saved and visually checked without another order. Preserve that distinction when describing what was verified.')
h(b,'Packing imagery')
p(b,'Shipping photographs show two closed retail canisters, each individually protected in clear bubble wrap inside a carton. Treat that as the recorded visual standard; warehouse procedures and final packing checks must follow the actual fulfillment process.')

b=page('The creative production workflow','07  WORKING WITH THE BRAND','S02 | S05 | Recommended execution sequence')
table(b,['Step','Action','Required output'],[
 ['1 Intake','Identify the exact scent, channel, audience, deliverables and latest owner instruction.','A short brief with clear scope.'],['2 Evidence','Read BRAND_GUIDE.md, both owner references, product photos and exact copy.','Fact sheet and source paths.'],['3 Creative mapping','Assign a distinct action, crop, props, palette and correct line casting.','A purposeful concept and shot list.'],['4 Photography','Create the scene with natural anatomy, product scale, light and contact shadows.','A clean photographic master.'],['5 Product fidelity','Check or composite the real product, preserving every identifying detail.','Approved bottle and packaging integration.'],['6 Typesetting','Add copy with actual fonts after the visual is ready.','Editable type layers and final export.'],['7 Review','Inspect identity, claims, faces, crop, readability and technical output.','QA record with revision notes.'],['8 Handoff','Save sources, prompts, final files, copy, alt text and manifest.','A traceable production package.']
 ],[0.9,3.3,2.35])
h(b,'A useful brief template')
p(b,'Deliverable and size: [channel and format]. Scent and line: [exact name and label category]. Objective: [one viewer action]. Approved copy: [source and exact text]. Scene: [one action and prop set]. Palette: [brand anchor, scent accent, neutral]. Crop: [face-free framing]. References: [real file paths]. QA owner and deadline: [assigned details].')
p(b,'For generation, explicitly reserve clean space for later typography. Never ask the model to invent a logo, bottle label, ingredient list, headline, barcode or product claim.')

b=page('Quality control before handoff','07  WORKING WITH THE BRAND','S02 | S05 | S08 | S11')
h(b,'Creative and product checks')
bullets(b,['The exact scent name and label category are correct; the product has not changed shape, color, liquid level or packaging.','Every visible body or hand matches the product line. Current campaign work contains no faces or partial faces anywhere.','Hands, grips, shadows, reflections and product scale look natural. No duplicate accessories or malformed objects appear.','The scene feels warm, active, playful and different from the other scents. The product is immediately identifiable.','All claims come from approved copy or package evidence. Ingredients are transcribed per scent.','Real production fonts are loaded. Headline, offer condition, COD and CTA are legible at intended mobile size.'])
h(b,'Website and ordering checks')
bullets(b,['Confirm native price, availability, variant and exact quantities. Test one item and two or more items, including duplicates of one scent.','Check narrow mobile and desktop layouts in both visual modes; inspect overlays, gallery controls, keyboard use and horizontal overflow.','Confirm actual order completion separately from opening a popup. Mocked tests, local preview checks and real saved orders are different evidence.','Preserve source artwork and font locks. Test the theme package in a draft before release. Keep the rollback theme and app-setting record.'])
h(b,'Release record')
p(b,'Save the version, date, task owner, source files, outputs, exact copy, checks performed, remaining issues and approval state. “Exported,” “reviewed,” “approved,” “uploaded” and “published” are separate statuses. A file named final or a folder with a newer date is not proof of approval.')

b=page('A practical content operating rhythm','07  WORKING WITH THE BRAND','Recommended workflow adapted from S15; current campaign scope remains six scents')
p(b,'Use a repeatable rhythm that the team can actually maintain. The schedule below is a working recommendation, not a booked media plan or proof that these activities have happened.')
table(b,['Content slot','Purpose','Example'],[['Scent story','Help shoppers choose','Exact approved sentence with a label-led visual.'],['Personal moment','Make the brand feel part of daily life','Face-free bag, music, park or ceramics scene.'],['Product detail','Build confidence','Real cap, canister construction or readable packaging.'],['Shopping information','Reduce uncertainty','Shipping threshold, COD and how to place an order.'],['Community or discovery','Invite participation','A scent-choice question, Finder walkthrough or permission-cleared customer story.']],[1.35,1.8,3.4])
h(b,'Weekly review')
p(b,'Review what was posted, what reached people, what drew useful questions and what produced valid orders. Separate creative attention from commercial results. For COD, distinguish placed, confirmed, fulfilled, canceled and returned orders so an attractive order count does not hide delivery losses.')
h(b,'Measurement definitions')
p(b,'Track spend, impressions, link clicks, product visits, placed orders, fulfilled orders, revenue, average order value, items per order and the share of orders with two or more items. Use a consistent attribution window and source. Do not claim ROAS, customer acquisition cost or revenue growth without the supporting records.')
h(b,'Creators and permissions')
p(b,'Give collaborators the approved scent copy, product facts, line-specific casting rules, face-free requirement and offer. Agree on deliverables, usage rights and disclosures before commercial reuse. A new creator contract, affiliate rate, giveaway or ad spend needs its own business authorization.')

b=page('Historical plans and active decisions','07  WORKING WITH THE BRAND','S15 Historical research | S02 | S09 | S11')
p(b,'The early market-research document explains the ambition behind the brand: bold visual identity, accessible confidence, scent-led self-expression and community content. Its numeric plans remain historical assumptions, not results or current spending instructions.')
table(b,['Historical planning item','Recorded figure','How to treat it'],[['90-day units target','6,000 units','A target; no achievement is established here.'],['90-day revenue target','About ₱4.7 million','At ₱799 × 6,000, gross product sales would be ₱4,794,000 before deductions.'],['Follower target','20,000','A target, not a verified audience count.'],['Monthly marketing budget','Above ₱100,000','Historical planning input; reconfirm before spending.'],['Shoot allocation','₱30,000–₱50,000','Historical estimate, not a vendor commitment.'],['Unit COGS assumption','₱170','Internal planning assumption; verify landed cost and all deductions.']],[1.65,1.5,3.4])
h(b,'What supersedes the early plan')
p(b,'Shopify is now the project’s master channel; the older TikTok Shop and Shopee plan does not prove those stores are launched. The active offer is free shipping on two or more items with COD; old ₱1,499 bundles and ₱50 vouchers are not active authority. The latest campaign is face-free; old portrait scripts are superseded.')
h(b,'Decisions to maintain')
p(b,'Assign owners for brand approval, product facts, content, website, customer care and fulfillment. Confirm current prices, stock, operational contact details, final policies, budgets and channel accounts before a release. Keep any personal founder story and unverified company history outside public copy until the owner provides it.')
p(b,'This page contains internal business planning. Share it only with collaborators who need those details; use the public-safe AI handoff on the next page for general creative work.','Caption')

b=page('Ready to paste AI handoff','08  AI AND COLLABORATOR HANDOFF','Current rules synthesized from S01–S11')
p(b,'Copy the text below into a new AI conversation, then attach the relevant reference images, product photographs and approved copy. A text prompt cannot substitute for visual product evidence. This version omits private budgets, costs, credentials and customer information.')
p(b,'You are working on DearBody Philippines. Use this playbook as the brand context and follow the owner’s latest instructions. The brand position is London-formulated fragrance culture curated for Filipino everyday life: premium but approachable, warm, vibrant, confident and playful. Tagline: A Scent Journey.')
p(b,'Official colors: burgundy #5C0006, red #9A1106, orange #D46601, golden tan #E9A250 and cream #F4E3CB. Use actual Cenzo Flare Bold for display copy and Helvetica Now for support. The recorded website body face is Helvetica Now Display Light. Use official logo assets without redrawing them.')
p(b,'Before any Dear Body image creation or edit, read BRAND_GUIDE.md and inspect both owner references: 2026-09-25-publication-materials-owner-reference.png and 2026-09-25-social-media-look-owner-reference.png. Adapt their playful warm saturated color, tactile photography, tight editorial crops and energetic objects. Do not copy them exactly.')
p(b,'Current Facebook scope: six ads, one per priority scent, plus profile and cover. All campaign lifestyle images and scripts are face-free, including mirrors, backgrounds, screens and printed portraits. Use Filipino hands, arms and torso-only crops. Women’s white or cream labels: Mojito Metallique, Amber Oud Silk, Mistened Narcissus. Men’s black labels: Charme Envoûtant, Oud Mirage, Rtulle & Satin. Match all visible talent to the product line.')
p(b,'Preserve exact packaging, spelling, proportions, cap, label and liquid. Use real product photographs as authority. Use only the approved scent descriptions and verified facts; do not invent notes, ingredients, benefits, performance, testimonials or offers. Typeset new copy after the photograph with real fonts.')
p(b,'Offer hierarchy: FREE SHIPPING ON 2+ ITEMS first, CASH ON DELIVERY second, and SHOP NOW. Do not expand the six-scent campaign to fourteen. Do not modify existing product galleries unless requested. The current website source is output/website/scent-journey-theme/theme. Use native Shopify commerce data; the local preview cannot place orders.')
p(b,'State the scope and source version, flag conflicts with newer owner directions, and return editable sources, final exports, exact copy and a concise QA record. Keep historical assets and plans distinct from current rules. If essential product evidence is absent, request it rather than inventing it. The specific task is: [insert the requested deliverable].')

b=page('The AI fact pack','08  AI AND COLLABORATOR HANDOFF','S02 | S03 | S06 | S09 | S10')
p(b,'These compact fields support reliable retrieval in other platforms. Preserve exact scent spelling and distinguish the display name from the source folder and Shopify handle.')
table(b,['Field','Value'],[['Brand','DearBody Philippines / Dear Body'],['Tagline','A Scent Journey'],['Identity source','Canva DEARBODY PLAYBOOK, DAHL_W_c7U8'],['Scent-copy source','Dear Body Brand Deep Dive, DAHIB2UjgJM'],['Priority count','6'],['Campaign assets','6 ads + 1 profile + 1 cover'],['Campaign casting','Face-free Filipino talent; line-matched hands, arms and torso'],['Display face','Cenzo Flare Bold'],['Website body face','Helvetica Now Display Light, weight 300'],['Current theme source','output/website/scent-journey-theme/theme/'],['Storefront','dearbody.ph'],['Order service','Shopify data + configured EasySell COD popup'],['Existing gallery handling','Preserve unless an edit is explicitly requested']],[2.0,4.55])
table(b,['Display name','Shopify handle'],[[s[1], 'rtulle-satin' if s[0]=='rtulle-and-satin' else s[0]] for s in SCENTS],[2.85,3.7])
p(b,'Do not place passwords, API keys, customer addresses, order tokens or confidential financial data into a public brand prompt. Upload only the materials needed for the task and make the latest owner instructions explicit.')

b=page('Project map and source authority','08  AI AND COLLABORATOR HANDOFF','Project-relative locations for portable handoff')
table(b,['Source','Use'],[
 ['S01  BRAND_GUIDE.md','Brand essence, mission, palette, typography, imagery and current social direction.'],
 ['S02  AGENTS.md','25 September owner constraints and campaign scope.'],
 ['S03  brand.tokens.json and brand.css','Digital identity values and implementation starting points.'],
 ['S04  handoff/assets/brand-references/','Both dated owner references and historical research PDFs.'],
 ['S05  PRODUCT_LISTING_IMAGE_STRATEGY.md','Product fidelity, intake, production and gallery standards; read later order revisions too.'],
 ['S06  output/product-listing/SCENT_DESCRIPTION_SOURCE_LEDGER.md','Exact approved description for each of the fourteen archived scents.'],
 ['S07  output/product-listing/<scent>/','SOURCE-FACTS, gallery manifests, final assets and scent-specific QA.'],
 ['S08  output/product-listing/WHO_IT_FITS_PERSONALITY_REVIEW_2026-09-24.md','Current personality copy. Pair with SHOPIFY_GALLERY_ORDER_2026-09-24.md.'],
 ['S09  output/website/scent-journey-theme/START-HERE.md','Current documented theme, pages, offer, Finder and asset behavior.'],
 ['S10  output/website/scent-journey-theme/source-assets/fonts/','Original font packages and provenance.'],
 ['S11  output/website/scent-journey-theme/UPDATE-COD-ORDERS.md','Published release record and same-page order verification.'],
 ['S12  output/facebook-ads/ and output/facebook-page-kit/','Existing social exports and sources; older faces and scripts require revision.'],
 ['S13  output/website/upload-ready/shopify-gallery-order-2026-09-24/','Fourteen ordered gallery packages with manifests.'],
 ['S14  output/website/scent-journey-theme/reviewed-product-galleries/','Six-scent website import and derivative provenance.'],
 ['S15  handoff/assets/brand-references/dear-body-edp-marketing-research-2026.pdf','Historical strategy and planning assumptions; not current claim or budget authority.']
 ],[3.5,3.05])

b=page('Keeping the playbook current','08  AI AND COLLABORATOR HANDOFF','Recommended governance for this master document')
h(b,'Resolve conflicts in this order')
p(b,'Use the owner’s latest explicit instruction for the current task first. Then apply current AGENTS.md and BRAND_GUIDE.md, verified product photographs and the approved copy ledger. For implementation details, consult the relevant current release source. Older prompts, research proposals and archived outputs remain historical evidence.')
h(b,'Known differences to preserve')
bullets(b,['The 25 September face-free rule replaces older portrait concepts for the current Facebook campaign and scripts. It does not silently rewrite existing gallery files.','The original production frame numbers differ from the later Shopify upload order. Use the current mapping and manifest.','The website’s Display Light font selection differs from older Text Regular exports. Locked raster assets retain their original typography.','The latest owner personality review supersedes old “For someone drawn to…” source-fact wording. Approved scent-description sentences stay exact.','Same-page EasySell COD replaces the older native-checkout redirect flow as the primary recorded ordering journey.','The fourteen-scent library is broader than the six-scent campaign.','Earlier research includes unapproved performance suggestions and historical budgets. Do not carry them into public copy or spending decisions.'])
h(b,'Revision process')
p(b,'When a rule changes, record the date, owner decision, affected scents, channels and files, and the evidence for the change. Update this Word master, the relevant source document and any AI briefing pack together. Keep the previous version in an archive and state whether existing assets need revision.')
h(b,'Version record')
table(b,['Version','Date','Scope'],[['1.0','25 September 2026','First consolidated master playbook covering the documented brand build, identity, product library, content, website, operations and AI handoff.']],[0.8,1.45,4.3])
p(b,'The brand is most consistent when the same verified facts, visual references and owner decisions travel with every new task. Use this document as the common starting point and keep its source trail intact.')

# Author an editable Word document with stable manual chapter pages.
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.60)
sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
sec.footer_distance=Inches(.28)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Bullet']:
    s=styles[name]; s.font.name='Helvetica Now Text'; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7)
    s.paragraph_format.line_spacing=Pt(14)
    s.font.size=Pt(10)
for name,size in [('Title',32),('Heading 1',25),('Heading 2',13)]:
    s=styles[name]; s.font.name='Cenzo Flare'; s.font.bold=True; s.font.size=Pt(size)
    s.paragraph_format.line_spacing=Pt(size*1.13); s.paragraph_format.space_after=Pt(10)
    s.paragraph_format.space_before=Pt(10 if name=='Heading 2' else 0)
    s.paragraph_format.keep_with_next=True
styles['Subtitle'].font.size=Pt(12); styles['Subtitle'].paragraph_format.line_spacing=Pt(17)
styles['Caption'].font.size=Pt(8); styles['Caption'].font.color.rgb=RGBColor.from_string('535353'); styles['Caption'].paragraph_format.line_spacing=Pt(10)
styles['List Bullet'].paragraph_format.left_indent=Inches(.16)
styles['List Bullet'].paragraph_format.first_line_indent=Inches(-.12)
for name,font,size in [('Specimen Display','Cenzo Flare',24),('Specimen Body','Helvetica Now Display',16),('Lead','Helvetica Now Text',12)]:
    s=styles.add_style(name,1); s.font.name=font; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
    s.font.bold=(name=='Specimen Display'); s.paragraph_format.line_spacing=Pt(size*1.2); s.paragraph_format.space_after=Pt(10)
styles['Normal'].paragraph_format.widow_control=True
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('DEAR BODY  ·  MASTER PLAYBOOK  ·  '); r.font.name='Helvetica Now Text'; r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

image_manifest=[]
def add_image(par,path,width,caption):
    data=(ROOT/path).read_bytes()
    # Store presentation copies at a practical print resolution. Source assets stay untouched.
    im=Image.open(io.BytesIO(data)); im.thumbnail((1800,1800))
    stream=io.BytesIO()
    if im.mode in ('RGBA','LA') or 'transparency' in im.info: im.save(stream,'PNG')
    else: im.convert('RGB').save(stream,'JPEG',quality=92,subsampling=0)
    stream.seek(0)
    r=par.add_run(); shape=r.add_picture(stream,width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    image_manifest.append({'source':path,'caption':caption,'original_size':Image.open(io.BytesIO(data)).size})

def add_table(block):
    t=doc.add_table(rows=1,cols=len(block['headers'])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=block['widths'] or [6.55/len(block['headers'])]*len(block['headers'])
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    rows=[block['headers']]+block['rows']
    for i,vals in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,txt) in enumerate(zip(cells,vals)):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                node=OxmlElement('w:'+side); node.set(qn('w:w'),'75'); node.set(qn('w:type'),'dxa'); mar.append(node)
            pr.append(mar)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'353535' if i==0 else ('F4F4F4' if i%2==0 else 'FFFFFF')); pr.append(shade)
            cp=cell.paragraphs[0]; cp.paragraph_format.space_after=Pt(0); cp.paragraph_format.line_spacing=Pt(11.5)
            run=cp.add_run(str(txt)); run.font.size=Pt(8.7); run.font.name='Helvetica Now Text'
            if i==0: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
        trpr=t.rows[i]._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); trpr.append(no)
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    aft=doc.add_paragraph(); aft.paragraph_format.space_after=Pt(1); aft.paragraph_format.line_spacing=Pt(4); aft.add_run().font.size=Pt(1)

def bookmark(par,name,id):
    a=OxmlElement('w:bookmarkStart'); a.set(qn('w:id'),str(id)); a.set(qn('w:name'),name)
    z=OxmlElement('w:bookmarkEnd'); z.set(qn('w:id'),str(id)); par._p.insert(0,a); par._p.append(z)

TOC=[('Brand in one page',3),('How the brand system was built',4),('Positioning and brand voice',5),('Color typography and visual references',7),('Photography and product fidelity',10),('Six priority scent profiles',11),('Wider library and gallery architecture',17),('Copy accuracy and claims',20),('Facebook concepts and video scripts',21),('Shopify website and visual examples',25),('Pages Finder pricing and ordering',27),('Delivery and customer care',30),('Production workflow and quality checks',31),('Content rhythm and historical plans',33),('AI handoff and fact pack',35),('Project source map and maintenance',37)]

for idx,pg in enumerate(pages):
    if idx: doc.add_page_break()
    ep=doc.add_paragraph(pg['section'],'Caption'); ep.paragraph_format.space_after=Pt(10)
    title=doc.add_paragraph(pg['title'],'Title' if idx==0 else 'Heading 1'); bookmark(title,'page_'+str(idx+1),idx+1)
    for block in pg['blocks']:
        k=block['kind']
        if k=='p': doc.add_paragraph(block['text'],block['style'])
        elif k=='table': add_table(block)
        elif k=='image':
            par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.space_after=Pt(4); par.paragraph_format.line_spacing=1
            add_image(par,block['path'],block['width'],block['caption']); par.paragraph_format.keep_with_next=True
            doc.add_paragraph(block['caption'],'Caption')
        elif k=='images':
            par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.line_spacing=1; par.paragraph_format.space_after=Pt(4); par.paragraph_format.keep_with_next=True
            for j,path in enumerate(block['paths']):
                if j: par.add_run('   ')
                add_image(par,path,block['width'],block['captions'][j])
            doc.add_paragraph('   |   '.join(block['captions']),'Caption')
        elif k=='palette':
            t=doc.add_table(rows=1,cols=5); t.autofit=False
            for cell,color in zip(t.rows[0].cells,['5C0006','9A1106','D46601','E9A250','F4E3CB']):
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),color); cell._tc.get_or_add_tcPr().append(sh)
                cell.paragraphs[0].paragraph_format.space_after=Pt(24); cell.paragraphs[0].add_run(' ')
            doc.add_paragraph('The five official digital colors','Caption')
        elif k=='toc':
            for label,num in TOC:
                par=doc.add_paragraph()
                par.paragraph_format.space_after=Pt(8)
                link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),'page_'+str(num))
                r=OxmlElement('w:r'); tx=OxmlElement('w:t'); tx.text=f'{label}   {num}'; r.append(tx); link.append(r); par._p.append(link)
    if pg['source']:
        sp=doc.add_paragraph(pg['source'],'Caption'); sp.paragraph_format.space_before=Pt(8)

doc.core_properties.title='Dear Body Master Playbook'
doc.core_properties.subject='Brand identity, products, creative production, website, operations and AI handoff'
doc.core_properties.author='DearBody Philippines'
doc.core_properties.keywords='Dear Body, brand playbook, brand guide, Shopify, fragrance, AI handoff'
doc.core_properties.comments='Version 1.0, 25 September 2026'
outfile=OUT/'Dear-Body-Master-Playbook.docx'
doc.save(outfile)

# Embed the supplied fonts using OOXML obfuscated font parts. All three report fsType 0.
from lxml import etree
NS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
OR='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
with zipfile.ZipFile(outfile) as z: data={n:z.read(n) for n in z.namelist()}
ft=etree.fromstring(data['word/fontTable.xml'])
rels=etree.Element('{'+REL+'}Relationships',nsmap={None:REL})
for n,(name,path) in enumerate(FONTS.items(),1):
    font=next((x for x in ft if x.get('{'+NS+'}name')==name),None)
    if font is None: font=etree.SubElement(ft,'{'+NS+'}font'); font.set('{'+NS+'}name',name)
    key=uuid.uuid4(); keybytes=bytes.fromhex(key.hex)[::-1]; raw=bytearray(path.read_bytes())
    for i in range(32): raw[i]^=keybytes[i%16]
    part=f'fonts/font{n}.odttf'; data['word/'+part]=bytes(raw)
    el=etree.SubElement(font,'{'+NS+'}embedBold' if name=='Cenzo Flare' else '{'+NS+'}embedRegular')
    el.set('{'+OR+'}id',f'rIdFont{n}'); el.set('{'+NS+'}fontKey','{'+str(key).upper()+'}')
    rr=etree.SubElement(rels,'{'+REL+'}Relationship'); rr.set('Id',f'rIdFont{n}'); rr.set('Type',OR+'/font'); rr.set('Target',part)
data['word/fontTable.xml']=etree.tostring(ft,xml_declaration=True,encoding='UTF-8',standalone=True)
data['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(data['[Content_Types].xml'])
el=etree.SubElement(ct,'{'+CT+'}Default'); el.set('Extension','odttf'); el.set('ContentType','application/vnd.openxmlformats-officedocument.obfuscatedFont')
data['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
st=etree.fromstring(data['word/settings.xml'])
etree.SubElement(st,'{'+NS+'}embedTrueTypeFonts')
data['word/settings.xml']=etree.tostring(st,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(outfile,'w',zipfile.ZIP_DEFLATED) as z:
    for name,body in data.items(): z.writestr(name,body)

(QA/'pages.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'image-sources.json').write_text(json.dumps(image_manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(outfile),'planned_pages':len(pages),'image_placements':len(image_manifest),'unique_images':len(set(x['source'] for x in image_manifest)),'titles':[p['title'] for p in pages]},ensure_ascii=False))
